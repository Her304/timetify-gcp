"""Validation tests for the syllabus-parsing schema and its failure paths.

These cover the contract between the LLM and the backend: what the model may
return, what is rejected, and how ambiguous input is represented. They run
entirely offline — no API key, no network — because the whole point of the
Pydantic layer is that correctness is checkable without the model.
"""

import os
import tempfile
from unittest import mock

from django.test import SimpleTestCase
from docx import Document
from pydantic import ValidationError

from main.pdf import (
    ExtractedCoursesResponse,
    _parse_spur_json,
    extract_text,
    process_course_outline,
)


def make_payload(**overrides):
    """A minimal, fully valid extraction payload."""
    course = {
        "course_id": "CS101",
        "course_name": "Intro to Computer Science",
        "term": "Spring 2026",
        "classroom": "Hall A",
        "start_time": "13:30",
        "end_time": "15:30",
        "rep_date": "Tuesday,Thursday",
        "is_main": True,
        "is_lab": False,
        "weeks": [
            {"week_number": 1, "week_date": "2026-05-04", "week_topic": "Overview"},
        ],
        "exams": [
            {"exam_date": "2026-06-15", "exam_topic": "Midterm"},
        ],
        "assignments": [
            {"assignment_due": "2026-05-11", "assignment_topic": "HW1"},
        ],
    }
    course.update(overrides.pop("course", {}))
    payload = {"courses": [course]}
    payload.update(overrides)
    return payload


class SchemaAcceptanceTests(SimpleTestCase):
    """What a well-formed model response looks like."""

    def test_valid_payload_parses_and_types_dates(self):
        parsed = ExtractedCoursesResponse.model_validate(make_payload())
        course = parsed.courses[0]
        self.assertTrue(course.is_main)
        self.assertEqual(course.weeks[0].week_topic, "Overview")
        # Typed date fields mean downstream code gets real `date` objects,
        # not strings that might be "May 4" or "TBD".
        self.assertEqual(course.exams[0].exam_date.isoformat(), "2026-06-15")

    def test_missing_information_is_null_not_fabricated(self):
        # The syllabus that says nothing about times/rooms: every optional
        # field is null and that must be accepted — an honest null is the
        # intended behaviour, not an error.
        payload = make_payload(course={
            "term": None, "classroom": None, "start_time": None,
            "end_time": None, "rep_date": None,
        })
        parsed = ExtractedCoursesResponse.model_validate(payload)
        self.assertIsNone(parsed.courses[0].classroom)
        self.assertIsNone(parsed.courses[0].rep_date)

    def test_labelled_guess_is_separate_from_stated_date(self):
        # "Final scheduled by registrar between Jul 31 and Aug 14": the exam
        # date itself must stay null; the model may attach a labelled guess.
        payload = make_payload(course={
            "exams": [{
                "exam_date": None,
                "exam_topic": "Final",
                "suggested_date": "2026-07-31",
                "suggestion_basis": "first day of exam window",
            }],
        })
        parsed = ExtractedCoursesResponse.model_validate(payload)
        exam = parsed.courses[0].exams[0]
        self.assertIsNone(exam.exam_date)
        self.assertEqual(exam.suggested_date.isoformat(), "2026-07-31")
        self.assertEqual(exam.suggestion_basis, "first day of exam window")


class SchemaRejectionTests(SimpleTestCase):
    """Ambiguous or malformed values must be rejected at the boundary."""

    def test_date_range_is_rejected(self):
        payload = make_payload(course={
            "exams": [{"exam_date": "2026-07-31 to 2026-08-14", "exam_topic": "Final"}],
        })
        with self.assertRaises(ValidationError):
            ExtractedCoursesResponse.model_validate(payload)

    def test_tbd_is_rejected(self):
        payload = make_payload(course={
            "assignments": [{"assignment_due": "TBD", "assignment_topic": "Essay"}],
        })
        with self.assertRaises(ValidationError):
            ExtractedCoursesResponse.model_validate(payload)

    def test_non_iso_date_is_rejected(self):
        payload = make_payload(course={
            "exams": [{"exam_date": "June 15", "exam_topic": "Midterm"}],
        })
        with self.assertRaises(ValidationError):
            ExtractedCoursesResponse.model_validate(payload)

    def test_missing_required_field_is_rejected(self):
        payload = make_payload()
        del payload["courses"][0]["course_id"]
        with self.assertRaises(ValidationError):
            ExtractedCoursesResponse.model_validate(payload)

    def test_wrong_item_shape_is_rejected(self):
        # e.g. the model collapsed an assignment table into one string.
        payload = make_payload(course={"assignments": ["weekly essays"]})
        with self.assertRaises(ValidationError):
            ExtractedCoursesResponse.model_validate(payload)


class RecurringAssignmentTests(SimpleTestCase):
    def test_weekly_recurrence_without_due_date_is_valid(self):
        payload = make_payload(course={
            "assignments": [{
                "assignment_due": None,
                "assignment_topic": "Online homework",
                "recurrence": "weekly",
                "recurrence_weekday": "Sunday",
            }],
        })
        parsed = ExtractedCoursesResponse.model_validate(payload)
        assignment = parsed.courses[0].assignments[0]
        self.assertEqual(assignment.recurrence, "weekly")
        self.assertEqual(assignment.recurrence_weekday, "Sunday")
        self.assertIsNone(assignment.assignment_due)


class SpurResponseParsingTests(SimpleTestCase):
    """The alternative provider path validates against the same schema."""

    def test_fenced_json_is_parsed(self):
        content = (
            "```json\n" + ExtractedCoursesResponse.model_validate(
                make_payload()).model_dump_json() + "\n```"
        )
        parsed = _parse_spur_json(content)
        self.assertEqual(parsed.courses[0].course_id, "CS101")

    def test_schema_mismatch_raises_value_error(self):
        with self.assertRaisesMessage(ValueError, "does not match the course schema"):
            _parse_spur_json('{"courses": [{"nonsense": true}]}')


class ProviderFailureTests(SimpleTestCase):
    """Configuration and provider errors surface as ValueError, not crashes."""

    def test_unknown_provider_is_rejected(self):
        env = {"COURSE_PARSER_PROVIDER": "banana", "OPENAI_API_KEY": "sk-test-dummy"}
        with mock.patch.dict(os.environ, env):
            with self.assertRaisesMessage(ValueError, "must be 'openai' or 'spur'"):
                process_course_outline("/nonexistent.pdf")

    def test_spur_without_api_key_is_rejected(self):
        env = {"COURSE_PARSER_PROVIDER": "spur", "SPUR_API_KEY": ""}
        with mock.patch.dict(os.environ, env):
            with self.assertRaisesMessage(ValueError, "SPUR_API_KEY is required"):
                process_course_outline("/nonexistent.pdf")


class TextExtractionTests(SimpleTestCase):
    def test_unsupported_extension_is_rejected(self):
        with self.assertRaisesMessage(ValueError, "Unsupported file extension"):
            extract_text("/tmp/syllabus.txt")

    def test_docx_paragraphs_are_extracted(self):
        doc = Document()
        doc.add_paragraph("CS101 — Spring 2026")
        doc.add_paragraph("Tues/Thurs 1:30-3:30, Hall A")
        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            doc.save(tmp.name)
            text = extract_text(tmp.name)
        self.assertIn("CS101", text)
        self.assertIn("Tues/Thurs", text)
